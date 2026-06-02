import os
import re
import shutil
from pathlib import Path
from typing import List, Dict, Optional

try:
    from llama_index.core import (
        VectorStoreIndex,
        SimpleDirectoryReader,
        StorageContext,
        load_index_from_storage,
        Settings,
        Document,
    )
    LLAMA_AVAILABLE = True
except ImportError:
    LLAMA_AVAILABLE = False

try:
    from llama_index.embeddings.huggingface import HuggingFaceEmbedding
    HF_AVAILABLE = True
except ImportError:
    HF_AVAILABLE = False

try:
    from llama_index.embeddings.openai import OpenAIEmbedding
    OPENAI_EMBED_AVAILABLE = True
except ImportError:
    OPENAI_EMBED_AVAILABLE = False

from .config import config


class VectorMemory:
    """
    Semantic long-term memory via LlamaIndex.
    Embedding priority: HuggingFace (local) → OpenAI → keyword mock
    """

    MAX_CHUNK = 500

    def __init__(self):
        self._index: Optional[object] = None
        self._available = False
        self._mock_store: List[Dict] = []

        if LLAMA_AVAILABLE:
            self._setup()

    # ── Public API ────────────────────────────────────────────

    def query(self, text: str, top_k: int = None) -> List[Dict]:
        k = top_k or config.TOP_K_VECTOR
        if not self._available:
            return self._mock_query(text, k)
        try:
            retriever = self._index.as_retriever(similarity_top_k=k)
            nodes = retriever.retrieve(text)
            return [
                {
                    "text":     node.text[: self.MAX_CHUNK],
                    "score":    round(node.score or 0.0, 3),
                    "source":   node.metadata.get("source", "vector"),
                    "filename": node.metadata.get("filename", ""),
                }
                for node in nodes
                if (node.score or 0) > 0.05
            ]
        except Exception as e:
            print(f"  [Vector] query error: {e}")
            return []

    def store(self, text: str, metadata: Dict = None):
        if not self._available:
            self._mock_store.append({"text": text, "metadata": metadata or {}})
            return
        try:
            doc = Document(text=text, metadata=metadata or {"source": "conversation"})
            self._index.insert(doc)
            self._index.storage_context.persist(persist_dir=config.INDEX_PERSIST_DIR)
        except Exception as e:
            print(f"  [Vector] store error: {e}")

    def index_path(self, path: str, recursive: bool = True) -> dict:
        p = Path(path)
        if not p.exists():
            return {"indexed": 0, "skipped": 0, "error": f"Path not found: {path}"}

        supported = {".txt", ".md"}
        try:
            import pypdf; supported.add(".pdf")  # noqa
        except ImportError:
            pass
        try:
            import docx; supported.add(".docx")  # noqa
        except ImportError:
            pass

        if p.is_file():
            file_list = [p] if p.suffix.lower() in supported else []
        elif recursive:
            file_list = [f for f in p.rglob("*") if f.is_file() and f.suffix.lower() in supported]
        else:
            file_list = [f for f in p.iterdir() if f.is_file() and f.suffix.lower() in supported]

        indexed, skipped, indexed_files = 0, 0, []
        for f in file_list:
            try:
                text = self._read_file(f)
                if not text or len(text.strip()) < 20:
                    skipped += 1
                    continue
                self.store(text, {"source": "file", "filename": f.name, "path": str(f)})
                indexed += 1
                indexed_files.append(str(f))
            except Exception as e:
                print(f"  [Vector] skip {f.name}: {e}")
                skipped += 1

        return {"indexed": indexed, "skipped": skipped, "files": indexed_files[:50]}

    def _read_file(self, path: Path) -> str:
        ext = path.suffix.lower()
        if ext in (".txt", ".md"):
            return path.read_text(encoding="utf-8", errors="ignore")
        elif ext == ".pdf":
            import pypdf
            reader = pypdf.PdfReader(str(path))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        elif ext == ".docx":
            import docx
            doc = docx.Document(str(path))
            return "\n".join(para.text for para in doc.paragraphs)
        return ""

    @property
    def is_available(self) -> bool:
        return self._available

    # ── Internal ──────────────────────────────────────────────

    def _setup(self):
        model_tag = None
        try:
            if HF_AVAILABLE:
                model_dir = Path(config.INDEX_PERSIST_DIR) / "models"
                Settings.embed_model = HuggingFaceEmbedding(
                    model_name="all-MiniLM-L6-v2",
                    cache_folder=str(model_dir),
                )
                model_tag = "hf-minilm-l6-v2"
                print("  [Vector] HuggingFace embeddings: all-MiniLM-L6-v2")
            elif OPENAI_EMBED_AVAILABLE and config.OPENAI_API_KEY:
                Settings.embed_model = OpenAIEmbedding(
                    model=config.EMBEDDING_MODEL,
                    api_key=config.OPENAI_API_KEY,
                    api_base=(
                        config.OPENAI_BASE_URL
                        if config.OPENAI_BASE_URL != "https://api.openai.com/v1"
                        else None
                    ),
                )
                model_tag = f"openai-{config.EMBEDDING_MODEL}"
                print(f"  [Vector] OpenAI embeddings: {config.EMBEDDING_MODEL}")
            else:
                raise Exception("No embedding provider available — install sentence-transformers")

            Settings.llm = None
            self._load_or_build_index(model_tag)
            self._available = True
        except Exception as e:
            print(f"  [Vector] setup failed: {e} — mock mode active")
            self._available = False

    def _load_or_build_index(self, model_tag: str):
        data_dir    = Path(config.DATA_DIR)
        persist_dir = Path(config.INDEX_PERSIST_DIR)
        data_dir.mkdir(parents=True, exist_ok=True)
        persist_dir.mkdir(parents=True, exist_ok=True)

        tag_file    = persist_dir / ".model_tag"
        stored_tag  = tag_file.read_text().strip() if tag_file.exists() else ""

        # Clear stale index if embedding model changed
        if stored_tag and stored_tag != model_tag and (persist_dir / "docstore.json").exists():
            print(f"  [Vector] Embedding model changed ({stored_tag} → {model_tag}), rebuilding index...")
            for f in persist_dir.glob("*.json"):
                f.unlink(missing_ok=True)

        # Load existing index
        if (persist_dir / "docstore.json").exists():
            try:
                ctx = StorageContext.from_defaults(persist_dir=str(persist_dir))
                self._index = load_index_from_storage(ctx)
                print(f"  [Vector] Loaded index ({model_tag})")
                return
            except Exception:
                pass

        # Build fresh index from /data files
        docs = []
        files = list(data_dir.rglob("*.txt")) + list(data_dir.rglob("*.md"))
        if files:
            docs = SimpleDirectoryReader(str(data_dir)).load_data()
        if not docs:
            docs = [Document(text="Memory initialized.", metadata={"source": "system"})]

        print(f"  [Vector] Building index from {len(docs)} document(s)...")
        self._index = VectorStoreIndex.from_documents(docs)
        self._index.storage_context.persist(persist_dir=str(persist_dir))
        tag_file.write_text(model_tag)
        print(f"  [Vector] Index built and saved ({model_tag})")

    # ── Mock fallback (keyword search, no embeddings) ─────────

    _MOCK_STOP = {
        "the","and","for","are","but","not","you","all","can","was","with",
        "this","that","have","from","they","will","been","were","your","just",
        "about","what","how","its","more","also","when","would","could","yes",
        "okay","yeah","did","does","got","get","now","see","let","i'm","do",
        "me","my","our","him","her","we","is","it","in","on","at","to","a",
        "of","be","so","if","as","or","by","an","no","up","he","she","was",
    }

    def _mock_query(self, text: str, k: int) -> List[Dict]:
        raw_words = set(re.sub(r"[^\w\s]", "", text.lower()).split())
        keywords = {w for w in raw_words if len(w) >= 4 and w not in self._MOCK_STOP}
        if not keywords:
            return []

        scored = []
        for item in self._mock_store[-200:]:
            item_words = set(re.sub(r"[^\w\s]", "", item["text"].lower()).split())
            overlap = keywords & item_words
            if len(overlap) >= 2 or (len(keywords) == 1 and overlap):
                score = len(overlap) / len(keywords)
                scored.append((score, item))

        scored.sort(key=lambda x: -x[0])
        return [
            {
                "text":     item["text"][: self.MAX_CHUNK],
                "score":    round(s, 3),
                "source":   item["metadata"].get("source", "mock-vector"),
                "filename": item["metadata"].get("filename", ""),
            }
            for s, item in scored[:k]
        ]
